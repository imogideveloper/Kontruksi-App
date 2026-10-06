# Copyright (c) 2026, Imogi Indonesia and contributors
# For license information, please see license.txt

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.model.naming import make_autoname
from frappe.utils import cstr, flt, getdate
from frappe.utils.number_format import NumberFormat
from frappe.utils.xlsxutils import read_xlsx_file_from_attached_file

# Kolom template Excel: (judul kolom, fieldname item).
KOLOM_EXCEL = (
	("Kode WBS", "kode_wbs"),
	("Uraian Pekerjaan", "uraian_pekerjaan"),
	("Spesifikasi / Keterangan", "spesifikasi"),
	("Satuan", "satuan"),
	("Volume", "volume"),
	("Harga Satuan (Rp)", "harga_satuan"),
	("Harga Satuan Pokok (Rp)", "harga_satuan_pokok"),
)

# Field item yang membentuk harga penawaran; dikunci setelah penawaran diajukan (biaya tetap boleh diubah).
FIELD_HARGA_ITEM = ("kode_wbs", "uraian_pekerjaan", "satuan", "volume", "harga_satuan")


class RABPenawaran(Document):
	def autoname(self):
		self.kode = make_autoname(f"RAB-{getdate().year}-.###", doc=self)
		self.name = self.kode

	def validate(self):
		if not self.is_new():
			self.kode = self.name
		self.cek_harga_terkunci()
		self.hitung_total()

	def onload(self):
		self.set_onload("harga_terkunci", penawaran_diajukan(self.tender))

	def cek_harga_terkunci(self):
		"""Setelah penawaran diajukan, harga yang dikirim ke panitia tidak boleh berubah; biaya pokok boleh."""
		if self.is_new() or not penawaran_diajukan(self.tender):
			return
		sebelum = self.get_doc_before_save()
		if not sebelum:
			return
		harga = lambda doc: [tuple(flt(item.get(f)) if f in ("volume", "harga_satuan") else cstr(item.get(f)) for f in FIELD_HARGA_ITEM) for item in doc.items]
		if harga(sebelum) != harga(self):
			frappe.throw(
				_("Penawaran tender {0} sudah diajukan, jadi item & harga RAB dikunci. Harga Satuan Pokok (biaya) tetap bisa diubah.").format(
					self.tender
				),
				title=_("Harga RAB dikunci"),
			)

	def hitung_total(self):
		tarif = flt(self.tarif_ppn) if self.status_ppn == "PPN" else 0

		for item in self.items:
			item.jumlah_harga = flt(flt(item.volume) * flt(item.harga_satuan), 2)
			item.ppn = flt(item.jumlah_harga * tarif / 100, 2)
			item.jumlah_harga_ppn = item.jumlah_harga + item.ppn
			item.jumlah_biaya = flt(flt(item.volume) * flt(item.harga_satuan_pokok), 2)
			item.margin_persen = (
				flt((item.jumlah_harga - item.jumlah_biaya) / item.jumlah_harga * 100, 2) if item.jumlah_harga and item.jumlah_biaya else 0
			)

		self.total_sebelum_ppn = sum(item.jumlah_harga for item in self.items)
		self.total_ppn = sum(item.ppn for item in self.items)
		self.total_rab = self.total_sebelum_ppn + self.total_ppn

		for item in self.items:
			item.bobot = flt(item.jumlah_harga / self.total_sebelum_ppn * 100, 2) if self.total_sebelum_ppn else 0

		# Persen HPS memakai Nilai Penawaran Kita dari Tender (angka yang diajukan), bukan total item RAB.
		self.persen_hps = flt(flt(self.nilai_penawaran) / flt(self.hps) * 100, 2) if flt(self.hps) else 0

		self.total_biaya = sum(flt(item.jumlah_biaya) for item in self.items)
		self.estimasi_margin = flt(self.total_sebelum_ppn) - flt(self.total_biaya) if self.total_biaya else 0
		self.persen_margin = (
			flt(self.estimasi_margin / self.total_sebelum_ppn * 100, 2) if self.total_biaya and self.total_sebelum_ppn else 0
		)


def penawaran_diajukan(tender):
	return bool(tender and frappe.db.get_value("Dokumen Tender", {"tender": tender}, "diajukan_pada"))


@frappe.whitelist()
def get_saran_uraian():
	"""Uraian pekerjaan yang pernah dipakai, untuk saran di kolom Uraian Pekerjaan."""
	return frappe.get_all(
		"RAB Penawaran Item",
		filters={"parenttype": "RAB Penawaran"},
		pluck="uraian_pekerjaan",
		distinct=True,
		order_by="uraian_pekerjaan",
		limit=500,
	)


def lihat_biaya():
	"""User boleh melihat kolom biaya (permission level 1 di RAB Penawaran)."""
	return 1 in frappe.get_meta("RAB Penawaran").get_permlevel_access("read")


def kolom_excel():
	return [k for k in KOLOM_EXCEL if k[1] != "harga_satuan_pokok" or lihat_biaya()]


def kirim_xlsx(nama_file, baris_item, catatan=None):
	"""File Excel RAB rapi: header berwarna, lebar & format angka, baris kelompok tebal, sheet Petunjuk."""
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

	kolom = kolom_excel()
	wb = Workbook()
	ws = wb.active
	ws.title = "RAB"
	garis = Border(bottom=Side(style="thin", color="D0D5DD"))
	lebar = {"kode_wbs": 11, "uraian_pekerjaan": 44, "spesifikasi": 34, "satuan": 9, "volume": 11,
		"harga_satuan": 19, "harga_satuan_pokok": 22}

	ws.append([judul for judul, _ in kolom])
	for c, (_, fieldname) in enumerate(kolom, start=1):
		sel = ws.cell(row=1, column=c)
		sel.font = Font(bold=True, color="2F5792")
		sel.fill = PatternFill("solid", fgColor="E3EFFF" if fieldname != "harga_satuan_pokok" else "FFF3D6")
		sel.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
		ws.column_dimensions[sel.column_letter].width = lebar.get(fieldname, 15)
	ws.row_dimensions[1].height = 30
	ws.freeze_panes = "A2"

	for item in baris_item:
		ws.append([item.get(fieldname) for _, fieldname in kolom])
		r = ws.max_row
		kelompok = cstr(item.get("kode_wbs")).strip() and "." not in cstr(item.get("kode_wbs"))
		for c, (_, fieldname) in enumerate(kolom, start=1):
			sel = ws.cell(row=r, column=c)
			sel.border = garis
			if fieldname in ("volume",):
				sel.number_format = "#,##0.##"
			elif fieldname in ("harga_satuan", "harga_satuan_pokok"):
				sel.number_format = "#,##0"
			if fieldname == "kode_wbs":
				sel.alignment = Alignment(horizontal="left")
			if kelompok:
				sel.font = Font(bold=True)
				sel.fill = PatternFill("solid", fgColor="F3F4F6")

	petunjuk = wb.create_sheet("Petunjuk")
	isi = [
		"Cara mengisi RAB Penawaran",
		"",
		"1. Isi sheet RAB mulai baris 2; jangan ubah judul kolom di baris 1.",
		"2. Kode WBS tanpa titik (1, 2, 3) = judul kelompok; kosongkan Satuan, Volume, dan Harga.",
		"3. Kode WBS bertitik (1.1, 1.2, 2.1) = item di bawah kelompoknya.",
		"4. Volume & harga diisi angka saja, tanpa 'Rp' (mis. 1250000 atau 12,5).",
	]
	if lihat_biaya():
		isi += [
			"5. Harga Satuan Pokok = biaya per satuan (material, upah, alat, subkon) tanpa keuntungan & PPN. Kolom ini internal.",
			"6. Upload: Excel > Upload Excel di form RAB. Pilih 'Perbarui Harga Pokok saja' untuk mengisi biaya",
			"   pada RAB yang sudah ada (dicocokkan lewat Kode WBS), juga setelah penawaran diajukan.",
		]
	else:
		isi.append("5. Upload: Excel > Upload Excel di form RAB.")
	if catatan:
		isi += ["", catatan]
	for baris in isi:
		petunjuk.append([baris])
	petunjuk["A1"].font = Font(bold=True, size=13)
	petunjuk.column_dimensions["A"].width = 110

	buffer = BytesIO()
	wb.save(buffer)
	frappe.response["filename"] = f"{nama_file}.xlsx"
	frappe.response["filecontent"] = buffer.getvalue()
	frappe.response["type"] = "binary"


@frappe.whitelist()
def download_template():
	frappe.has_permission("RAB Penawaran", "read", throw=True)
	contoh = [
		{"kode_wbs": "1", "uraian_pekerjaan": "Pekerjaan Persiapan"},
		{"kode_wbs": "1.1", "uraian_pekerjaan": "Pembersihan lokasi", "spesifikasi": "Termasuk buang puing", "satuan": "m2",
			"volume": 250, "harga_satuan": 15000, "harga_satuan_pokok": 11000},
		{"kode_wbs": "2", "uraian_pekerjaan": "Pekerjaan Beton"},
		{"kode_wbs": "2.1", "uraian_pekerjaan": "Beton K-250", "spesifikasi": "Ready mix", "satuan": "m3",
			"volume": 12.5, "harga_satuan": 1250000, "harga_satuan_pokok": 1050000},
	]
	kirim_xlsx("Template RAB Penawaran", contoh, catatan="Baris contoh di sheet RAB boleh dihapus / ditimpa.")


@frappe.whitelist()
def download_isi(name):
	"""Isi RAB yang sudah ada dalam format template, untuk diedit lalu diupload ulang."""
	doc = frappe.get_doc("RAB Penawaran", name)
	doc.check_permission("read")
	kirim_xlsx(f"{doc.name} - {doc.nama_project or ''}".strip(" -"), [item.as_dict() for item in doc.items])


@frappe.whitelist()
def baca_excel(file_url):
	"""Baca file Excel template RAB, kembalikan daftar item. File upload dihapus setelah dibaca."""
	frappe.has_permission("RAB Penawaran", "write", throw=True)
	rows = read_xlsx_file_from_attached_file(file_url=file_url)
	frappe.delete_doc("File", frappe.db.get_value("File", {"file_url": file_url}), ignore_permissions=True)

	if not rows:
		frappe.throw(_("File Excel kosong."))

	judul = [cstr(h).strip().lower() for h in rows[0]]
	posisi = {}
	for nama_kolom, fieldname in KOLOM_EXCEL:
		if nama_kolom.lower() in judul:
			posisi[fieldname] = judul.index(nama_kolom.lower())
	if "uraian_pekerjaan" not in posisi:
		frappe.throw(_("Kolom 'Uraian Pekerjaan' tidak ditemukan. Gunakan Download Template."))

	items = []
	for baris in rows[1:]:
		item = {f: (baris[i] if i < len(baris) else None) for f, i in posisi.items()}
		if not cstr(item.get("uraian_pekerjaan")).strip():
			continue
		for f in ("kode_wbs", "uraian_pekerjaan", "spesifikasi", "satuan"):
			if f in item:
				item[f] = cstr(item[f]).strip()
		for f in ("volume", "harga_satuan", "harga_satuan_pokok"):
			if f in item:
				item[f] = parse_angka(item[f])
		items.append(item)

	if not items:
		frappe.throw(_("Tidak ada baris item yang terisi di file Excel."))
	return items


def parse_angka(value):
	"""Angka dari Excel bisa berupa number atau teks berformat lokal (mis. 'Rp 1.250.000,50')."""
	if isinstance(value, int | float):
		return flt(value)
	teks = cstr(value).replace("Rp", "").replace(" ", "").strip()
	if not teks:
		return 0
	fmt = NumberFormat.from_string(frappe.db.get_default("number_format") or "#,###.##")
	return flt(teks.replace(fmt.thousands_separator, "").replace(fmt.decimal_separator, "."))


@frappe.whitelist()
def get_template_rab(tender):
	"""Template RAB dari master Jenis Project tender ini: {jenis_project, rows: [{kode_wbs, uraian_pekerjaan, satuan}]}."""
	frappe.has_permission("Tender", "read", tender, throw=True)
	jenis = frappe.db.get_value("Tender", tender, "jenis_project")
	if not jenis:
		frappe.throw(_("Jenis Project tender {0} belum diisi.").format(tender))
	rows = frappe.get_all("Jenis Project RAB", filters={"parent": jenis, "parenttype": "Jenis Project"},
		fields=["kode_wbs", "uraian_pekerjaan", "satuan"], order_by="idx asc")
	return {"jenis_project": jenis, "rows": rows}


FIELD_DARI_TENDER = {
	"nama_project": "nama_paket",
	"pemberi_kerja": "pemberi_kerja",
	"hps": "hps",
	"nilai_penawaran": "nilai_penawaran",
	"status_ppn": "status_ppn",
	"tarif_ppn": "tarif_ppn",
}


def sinkron_dari_tender(tender, method=None):
	"""Tender.on_update: salin Nama Project, Pemberi Kerja, HPS, Nilai Penawaran, Status & Tarif PPN ke RAB Penawaran tender itu, lalu
	hitung ulang total, PPN & persen HPS. db_update (bukan save) supaya kunci harga RAB tidak menggagalkan simpan Tender."""
	from frappe.utils import now

	from konstruksi.api import beri_tahu_form

	for name in frappe.get_all("RAB Penawaran", filters={"tender": tender.name, "docstatus": ("<", 2)}, pluck="name"):
		doc = frappe.get_doc("RAB Penawaran", name)
		berubah = False
		for field_rab, field_tender in FIELD_DARI_TENDER.items():
			if doc.get(field_rab) != tender.get(field_tender):
				doc.set(field_rab, tender.get(field_tender))
				berubah = True
		if not berubah:
			continue
		doc.hitung_total()
		doc.modified = now()
		doc.db_update_all()
		beri_tahu_form("RAB Penawaran", name)
